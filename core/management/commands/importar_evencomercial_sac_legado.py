from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import (
    AcaoEmEspera,
    Cliente,
    Empresa,
    ImportacaoLegadoSAC,
    ItemNotaFiscal,
    NotaFiscal,
    SAC,
    SACItem,
    SacHistorico,
    Setor,
    StatusSAC,
    TipoOcorrencia,
)
from core.services.texto_padrao import padronizar_caixa_texto


EMPRESA_CODIGO = "EVENLAB"
ANOS_IMPORTAVEIS = {2024, 2026}


ACAO_PADRAO_MAP = {
    "negociacao com cliente troca cancelamento": (
        "Aguardando Contato com Cliente - SAC",
        "Licitação",
    ),
    "aguardando faturamento expedicao": (
        "Aguardando Emissão da Nota fiscal (Saída)",
        "Logística",
    ),
    "sac emitir pedido cliente": (
        "Aguardando Emissão do Pedido (Saída)",
        "SAC",
    ),
    "em transito aguardando chegar no cliente": (
        "Em Transito (Com destino ao Cliente)",
        "Logística",
    ),
    "aguardando coleta reversa no cliente": (
        "Aguardando Coleta no Cliente",
        "Logística",
    ),
    "aguardando autorizacao diretoria": (
        "Aguardando Aprovação de orçamento (Interno Diretoria)",
        "Diretoria",
    ),
    "aguardando tecnico externo": (
        "Aguardando Visita Tecnica Externa",
        "Assistência Técnica",
    ),
    "aguardando manutencao": (
        "Aguardando Manutenção Interna",
        "Assistência Técnica",
    ),
    "aguardando logistica confirmar recebimento": (
        "Aguardando Recebimento Fisico/Fiscal",
        "Logística",
    ),
    "em transito aguardando chegar na ionlab": (
        "Em Transito (Com destino a Ionlab)",
        "Logística",
    ),
    "aguardando cliente confirmar recebimento aceite": (
        "Confirmar entrega realizada no cliente",
        "SAC",
    ),
    "aguardando orcamento tecnico externo": (
        "Aguardando Manutenção Externa",
        "Assistência Técnica",
    ),
}


TIPO_OCORRENCIA_MAP = {
    "material equipamento faltante": "Itens/Equipamentos faltantes na embalagem",
    "material equipamento errado": "Itens/Equipamentos em desacordo com pedido",
    "tecnico": "Tecnico - Equipamento não funciona",
    "avarias": "Itens/Equipamentos avariados",
    "material equipamento sobrando": "Itens/Equipamentos sobrando na embalagem",
}


SETOR_MAP = {
    "logistica": "Logística",
    "sac": "SAC",
    "assessoria cientifica": "Assessoria Cientifica",
    "licitacao": "Licitação",
    "financeiro": "Financeiro",
    "diretoria": "Diretoria",
    "tecnica": "Assistência Técnica",
    "assistencia tecnica": "Assistência Técnica",
    "producao": "Produção",
}


CRONOLOGIA_EVENTOS = [
    ("aberto", "Aberto", "SAC importado da planilha historica Even."),
    (
        "aberto aguardando tratativa tecnica por telefone",
        "Aberto - Aguardando Tratativa Tecnica Por Telefone",
        "Etapa historica: aguardando tratativa tecnica por telefone.",
    ),
    ("aberto aguardando chegada", "Aberto - Aguardando Chegada", "Etapa historica: aguardando chegada."),
    ("aberto aguardando avaliacao", "Aberto - Aguardando Avaliação", "Etapa historica: aguardando avaliacao."),
    (
        "aberto aguardando aprov orcamento",
        "Aberto - Aguardando Aprov. Orçamento",
        "Etapa historica: aguardando aprovacao de orcamento.",
    ),
    ("aberto aguardando peca", "Aberto - Aguardando Peça", "Etapa historica: aguardando peca."),
    (
        "aberto aguardando manutencao interna",
        "Aberto - Aguardando manutenção Interna",
        "Etapa historica: aguardando manutencao interna.",
    ),
    (
        "aberto aguardando manutencao externa",
        "Aberto - Aguardando manutenção Externa",
        "Etapa historica: aguardando manutencao externa.",
    ),
    (
        "aberto aguardando devolucao ao cliente",
        "Aberto - Aguardando devolução ao Cliente",
        "Etapa historica: aguardando devolucao ao cliente.",
    ),
    (
        "encerrado parcialmente aguardando entrega ao cliente",
        "Encerrado Parcialmente - Aguardando entrega ao cliente",
        "Etapa historica: encerrado parcialmente, aguardando entrega ao cliente.",
    ),
    ("acao interna", "Ação interna", "Etapa historica: acao interna."),
    ("encerrado", "Encerrado", "SAC concluido conforme planilha historica Even."),
]


def norm(valor):
    if pd.isna(valor):
        return ""
    texto = str(valor or "").strip().replace("\xa0", " ")
    return re.sub(r"\s+", " ", texto).strip()


def norm_ascii(valor):
    texto = unicodedata.normalize("NFKD", norm(valor)).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[^a-zA-Z0-9]+", " ", texto.lower())
    return re.sub(r"\s+", " ", texto).strip()


def numero_texto(valor):
    texto = norm(valor)
    if not texto:
        return ""
    if re.fullmatch(r"\d+\.0", texto):
        return texto[:-2]
    try:
        numero = float(texto.replace(",", "."))
        if numero.is_integer():
            return str(int(numero))
    except Exception:
        pass
    return texto


def inteiro(valor, padrao=1):
    texto = numero_texto(valor)
    if not texto:
        return padrao
    try:
        return max(1, int(Decimal(texto.replace(",", "."))))
    except (InvalidOperation, ValueError):
        return padrao


def data_valor(valor):
    if pd.isna(valor) or valor in ("", None):
        return None
    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()
    if isinstance(valor, datetime):
        return valor
    parsed = pd.to_datetime(valor, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.to_pydatetime()


def data_hora_aware(valor, fim_do_dia=False):
    dt = data_valor(valor)
    if not dt:
        return None
    if not isinstance(dt, datetime):
        dt = datetime.combine(dt, time.max if fim_do_dia else time.min)
    if dt.time() == time.min and fim_do_dia:
        dt = datetime.combine(dt.date(), time.max.replace(microsecond=0))
    if timezone.is_naive(dt):
        return timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def limitar(valor, tamanho):
    texto = norm(valor)
    return texto[:tamanho] if texto else ""


def coluna_por_chave(df, chave):
    for col in df.columns:
        if norm_ascii(col) == chave:
            return col
    for col in df.columns:
        if chave and chave in norm_ascii(col):
            return col
    return None


def valor_linha(row, col):
    if col is None or col not in row.index:
        return None
    return row.get(col)


def historico_formatado(texto):
    bruto = norm(texto)
    if not bruto:
        return ""
    bruto = re.sub(r"\*{2,}", "\n", bruto)
    bruto = re.sub(r"(?<!^)(?<!\n)(?<!\d)(\d{1,2}/\d{1,2}/\d{2,4})\s*[-–:]", r"\n\1 -", bruto)
    linhas = []
    atual = ""
    for parte in re.split(r"\n+", bruto):
        parte = parte.strip(" -*")
        if not parte:
            continue
        parte = re.sub(r"^(\d{1,2}/\d{1,2}/\d{2,4})\s*[-–:]?\s*", r"\1 - ", parte)
        parte = re.sub(r"\s+-\s+", " - ", parte)
        if re.match(r"^\d{1,2}/\d{1,2}/\d{2,4}\s+-\s+", parte):
            if atual:
                linhas.append(atual.strip())
            atual = parte
        elif atual:
            atual += " " + parte
        else:
            linhas.append(parte)
    if atual:
        linhas.append(atual.strip())
    return padronizar_caixa_texto("\n".join(linhas))


def sac_numero_sistema(numero_origem, ano):
    numero = int(numero_texto(numero_origem))
    return f"{numero:03d}/{ano}", numero


class ImportadorEvenComercialSAC:
    def __init__(self, *, arquivo, commit=False, log_output="", usuario=None):
        self.arquivo = Path(arquivo)
        self.commit = commit
        self.log_output = Path(log_output) if log_output else None
        self.usuario = usuario
        self.empresa = Empresa.objects.filter(codigo__iexact=EMPRESA_CODIGO).first()
        if not self.empresa:
            raise CommandError(f"Empresa {EMPRESA_CODIGO} nao encontrada.")

        self.status_aberto = self._status("Aberto")
        self.status_analise = self._status("Em Análise") or self.status_aberto
        self.status_concluido = self._status("Concluído")
        if not self.status_aberto or not self.status_concluido:
            raise CommandError("Status obrigatorios nao encontrados: Aberto/Concluido.")

        self.setores = {norm_ascii(s.nome): s for s in Setor.objects.filter(ativo=True)}
        self.acoes = {norm_ascii(a.nome): a for a in AcaoEmEspera.objects.filter(ativo=True)}
        self.tipos = {norm_ascii(t.nome): t for t in TipoOcorrencia.objects.filter(ativo=True)}
        self.stats = defaultdict(int)
        self.logs = []
        self.importados = []

    def _status(self, nome):
        chave = norm_ascii(nome)
        for status in StatusSAC.objects.filter(ativo=True):
            if norm_ascii(status.nome) == chave:
                return status
        return None

    def _log(self, severidade, sac, linha, campo, valor, motivo):
        self.logs.append(
            {
                "severidade": severidade,
                "sac": sac or "",
                "linha": linha or "",
                "campo": campo,
                "valor": limitar(valor, 500),
                "motivo": motivo,
            }
        )

    def _cols(self, df):
        cronologia = {}
        for chave, _, _ in CRONOLOGIA_EVENTOS:
            cronologia[chave] = coluna_por_chave(df, chave)
        return {
            "numero": coluna_por_chave(df, "no sac") or coluna_por_chave(df, "no"),
            "status": coluna_por_chave(df, "status"),
            "aberto": coluna_por_chave(df, "aberto"),
            "acao_padrao": coluna_por_chave(df, "acao padrao"),
            "agenda": coluna_por_chave(df, "agenda previsao"),
            "todo": coluna_por_chave(df, "to do"),
            "aguarda_acao": coluna_por_chave(df, "aguarda acao de"),
            "historico": coluna_por_chave(df, "qual acao pendente"),
            "nota": coluna_por_chave(df, "nota fiscal venda"),
            "emissao": coluna_por_chave(df, "emissao"),
            "cliente_codigo": coluna_por_chave(df, "cod cli"),
            "cliente": coluna_por_chave(df, "cliente"),
            "contato": coluna_por_chave(df, "contato"),
            "telefone": coluna_por_chave(df, "ddd telefone"),
            "email": coluna_por_chave(df, "e mail"),
            "referencia": coluna_por_chave(df, "referencia"),
            "equipamento": coluna_por_chave(df, "equipamento"),
            "quantidade": coluna_por_chave(df, "quantidade"),
            "serie": coluna_por_chave(df, "no serie"),
            "tempo_uso": coluna_por_chave(df, "tempo de uso"),
            "nf_cliente": coluna_por_chave(df, "nf do cliente"),
            "dt_emissao_cliente": coluna_por_chave(df, "dt emissao"),
            "relato": coluna_por_chave(df, "o que o cliente alegou de problema"),
            "parecer": coluna_por_chave(df, "resumo parecer area tecnica interna terceira"),
            "solucao": coluna_por_chave(df, "solucao padrao"),
            "tipo_2024": coluna_por_chave(df, "tipo de problema"),
            "tipo_2026": coluna_por_chave(df, "datas de cada acao"),
            "parte_problema": coluna_por_chave(df, "indique a parte do equipamento com problema"),
            "categoria": coluna_por_chave(df, "categoria"),
            "origem": coluna_por_chave(df, "setor da origem do problema"),
            "ato": coluna_por_chave(df, "ato que gerou o problema"),
            "garantia": coluna_por_chave(df, "dentro do periodo de garantia"),
            "mau_uso": coluna_por_chave(df, "detectado mau uso"),
            "conserto_garantia": coluna_por_chave(df, "conserto em garantia"),
            "custo_pecas": coluna_por_chave(df, "custo total das pecas de reposicao"),
            "frete_ion": coluna_por_chave(df, "valor frete r de envio para ion"),
            "frete_cliente": coluna_por_chave(df, "valor do frete de envio ao cliente"),
            "custo_tecnico": coluna_por_chave(df, "custo com tecnico terceiro"),
            "valor_manutencao": coluna_por_chave(df, "valor da manutencao cobrada do cliente"),
            "cronologia": cronologia,
        }

    def _acao_setor_por_padrao(self, valor):
        regra = ACAO_PADRAO_MAP.get(norm_ascii(valor))
        if not regra:
            return None, None
        acao_nome, setor_nome = regra
        return self.acoes.get(norm_ascii(acao_nome)), self.setores.get(norm_ascii(setor_nome))

    def _setor(self, valor):
        chave = norm_ascii(valor)
        if not chave or chave == "cliente":
            return None
        nome = SETOR_MAP.get(chave, valor)
        return self.setores.get(norm_ascii(nome))

    def _tipo_ocorrencia(self, valor):
        chave = norm_ascii(re.sub(r"^\d+\s*-\s*", "", norm(valor)))
        if not chave or chave == "administrativo":
            return None
        nome = TIPO_OCORRENCIA_MAP.get(chave, valor)
        return self.tipos.get(norm_ascii(nome))

    def _item_nf(self, nf, referencia, equipamento):
        qs = ItemNotaFiscal.objects.filter(nota_fiscal=nf)
        if referencia:
            item = qs.filter(codigo_produto__iexact=referencia).first()
            if item:
                return item
        if equipamento:
            return qs.filter(descricao_item__iexact=equipamento).first()
        return None

    def carregar(self):
        if not self.arquivo.exists():
            raise CommandError(f"Arquivo nao encontrado: {self.arquivo}")
        linhas = []
        for sheet in pd.ExcelFile(self.arquivo).sheet_names:
            df = pd.read_excel(self.arquivo, sheet_name=sheet, dtype=object)
            cols = self._cols(df)
            if not cols["numero"] or not cols["aberto"]:
                continue
            for idx, row in df.iterrows():
                numero = numero_texto(valor_linha(row, cols["numero"]))
                abertura = data_hora_aware(valor_linha(row, cols["aberto"]))
                if not numero or not abertura:
                    continue
                ano = abertura.year
                if ano == 2024 and sheet != "2024":
                    self._log("IGNORADO", f"{int(numero):03d}/2024", idx + 2, "Aba", sheet, "Duplicado em outra aba; fonte escolhida: 2024.")
                    continue
                if ano == 2026 and sheet != "2025 e 2026":
                    self._log("IGNORADO", f"{int(numero):03d}/2026", idx + 2, "Aba", sheet, "Fonte de 2026 deve ser a aba 2025 e 2026.")
                    continue
                if ano not in ANOS_IMPORTAVEIS:
                    self._log("IGNORADO", f"{int(numero):03d}/{ano}", idx + 2, "Ano", ano, "Ano fora desta rodada de importacao.")
                    continue
                linhas.append(
                    {
                        "sheet": sheet,
                        "linha": idx + 2,
                        "row": row,
                        "cols": cols,
                        "numero_origem": numero,
                        "ano": ano,
                        "abertura": abertura,
                        "chave": (numero, ano),
                    }
                )
        return linhas

    def _eventos_grupo(self, grupo):
        eventos = {}
        for item in grupo:
            row = item["row"]
            cols = item["cols"]
            for chave, label, obs in CRONOLOGIA_EVENTOS:
                col = cols["cronologia"].get(chave)
                data = data_hora_aware(valor_linha(row, col), fim_do_dia=(chave == "encerrado"))
                if not data:
                    continue
                atual = eventos.get(chave)
                if not atual or data > atual["data"]:
                    eventos[chave] = {"chave": chave, "label": label, "obs": obs, "data": data}
        return eventos

    def validar_grupo(self, numero_origem, ano, grupo):
        pendencias = []
        avisos = []
        primeira = grupo[0]
        row = primeira["row"]
        cols = primeira["cols"]
        numero_sac, _ = sac_numero_sistema(numero_origem, ano)

        status_chave = norm_ascii(valor_linha(row, cols["status"]))
        concluido = status_chave == "encerrado"

        nf_numero = numero_texto(valor_linha(row, cols["nota"]))
        nf = None
        if nf_numero:
            nf = (
                NotaFiscal.objects.filter(empresa=self.empresa, numero_nf=nf_numero).first()
                or NotaFiscal.objects.filter(empresa=self.empresa, numero=nf_numero).first()
            )
        if not nf_numero:
            pendencias.append("Nota fiscal nao informada.")
            self._log("ERRO", numero_sac, primeira["linha"], "Nota Fiscal Venda", "", "Nota fiscal nao informada.")
        elif not nf:
            pendencias.append(f"Nota fiscal nao encontrada: {nf_numero}.")
            self._log("ERRO", numero_sac, primeira["linha"], "Nota Fiscal Venda", nf_numero, "Nota fiscal nao encontrada na base da Even.")

        codigo_cliente = numero_texto(valor_linha(row, cols["cliente_codigo"]))
        cliente = Cliente.objects.filter(empresa=self.empresa, codigo_interno=codigo_cliente).first() if codigo_cliente else None
        if not cliente and nf:
            cliente = nf.cliente
        if codigo_cliente and nf and cliente and nf.cliente_id != cliente.id:
            avisos.append("Cliente da planilha difere do cliente da NF.")
            self._log("AVISO", numero_sac, primeira["linha"], "Cliente/NF", codigo_cliente, "Cliente da planilha difere do cliente vinculado a NF.")
        if not cliente:
            pendencias.append("Cliente nao encontrado.")
            self._log("ERRO", numero_sac, primeira["linha"], "Cliente", codigo_cliente, "Cliente nao localizado pela planilha nem pela NF.")

        tipo_valor = valor_linha(row, cols["tipo_2024"]) or valor_linha(row, cols["tipo_2026"])
        tipo = self._tipo_ocorrencia(tipo_valor)
        if norm_ascii(tipo_valor) == "administrativo":
            pendencias.append("Tipo Administrativo nao importado automaticamente.")
            self._log("ERRO", numero_sac, primeira["linha"], "Tipo de ocorrencia", tipo_valor, "Tipo Administrativo precisa de confirmacao/correcao.")
        elif not tipo:
            pendencias.append("Tipo de ocorrencia nao mapeado.")
            self._log("ERRO", numero_sac, primeira["linha"], "Tipo de ocorrencia", tipo_valor, "Tipo de ocorrencia nao encontrado no sistema.")

        acao, setor_regra = (None, None) if concluido else self._acao_setor_por_padrao(valor_linha(row, cols["acao_padrao"]))
        setor = None if concluido else (self._setor(valor_linha(row, cols["aguarda_acao"])) or setor_regra)
        if not concluido and not setor:
            pendencias.append("Setor atual nao mapeado.")
            self._log("ERRO", numero_sac, primeira["linha"], "Aguarda acao de", valor_linha(row, cols["aguarda_acao"]), "Setor atual nao mapeado.")
        if not concluido and not acao:
            pendencias.append("Acao em espera nao mapeada.")
            self._log("ERRO", numero_sac, primeira["linha"], "Acao Padrao", valor_linha(row, cols["acao_padrao"]), "Acao em espera nao mapeada.")

        itens_validos = 0
        for item in grupo:
            row_item = item["row"]
            cols_item = item["cols"]
            referencia = numero_texto(valor_linha(row_item, cols_item["referencia"]))
            equipamento = norm(valor_linha(row_item, cols_item["equipamento"]))
            categoria = norm_ascii(valor_linha(row_item, cols_item["categoria"]))
            rastreio = numero_texto(valor_linha(row_item, cols_item["serie"]))
            linha = item["linha"]

            if not referencia and not equipamento:
                pendencias.append(f"Linha {linha}: item/referencia nao informado.")
                self._log("ERRO", numero_sac, linha, "Item", "", "Item/referencia nao informado.")
                continue
            if categoria == "equipamento" and not rastreio:
                pendencias.append(f"Linha {linha}: equipamento sem numero de serie.")
                self._log("ERRO", numero_sac, linha, "No Serie", referencia or equipamento, "Equipamento sem numero de serie.")
            item_nf = self._item_nf(nf, referencia, equipamento) if nf else None
            if not item_nf:
                pendencias.append(f"Linha {linha}: item nao encontrado na NF {nf_numero}.")
                self._log("ERRO", numero_sac, linha, "Item da NF", referencia or equipamento, "Item nao encontrado na nota fiscal.")
            else:
                itens_validos += 1

        if not itens_validos:
            pendencias.append("Nenhum item valido localizado para o SAC.")

        return {
            "numero_sac": numero_sac,
            "cliente": cliente,
            "nota": nf,
            "tipo": tipo,
            "concluido": concluido,
            "acao": acao,
            "setor": setor,
            "pendencias": pendencias,
            "avisos": avisos,
        }

    def observacao_sac(self, row, cols):
        partes = []
        relato = norm(valor_linha(row, cols["relato"]))
        parecer = norm(valor_linha(row, cols["parecer"]))
        solucao = norm(valor_linha(row, cols["solucao"]))
        if relato:
            partes.append(f"Relato do cliente/usuario sobre o problema:\n{relato}")
        if parecer:
            partes.append(f"Resumo/Parecer tecnico interno ou terceiro:\n{parecer}")
        if solucao:
            partes.append(f"Solucao padrao:\n{solucao}")
        return padronizar_caixa_texto("\n\n".join(partes))

    def observacao_item(self, row, cols):
        campos = [
            ("Tipo de problema", cols["tipo_2024"] or cols["tipo_2026"]),
            ("Parte do equipamento", cols["parte_problema"]),
            ("Categoria", cols["categoria"]),
            ("Setor da origem do problema", cols["origem"]),
            ("Ato que gerou o problema", cols["ato"]),
            ("Dentro do periodo de garantia", cols["garantia"]),
            ("Detectado mau uso", cols["mau_uso"]),
            ("Conserto em garantia", cols["conserto_garantia"]),
            ("Custo total das pecas", cols["custo_pecas"]),
            ("Frete envio para Ion", cols["frete_ion"]),
            ("Frete envio ao cliente", cols["frete_cliente"]),
            ("Custo tecnico terceiro", cols["custo_tecnico"]),
            ("Valor manutencao cliente", cols["valor_manutencao"]),
        ]
        linhas = []
        for rotulo, col in campos:
            valor = norm(valor_linha(row, col))
            if valor:
                linhas.append(f"{rotulo}: {valor}")
        return padronizar_caixa_texto("\n".join(linhas))

    def _tempo_uso_dias(self, nf, data_abertura):
        data_nf = getattr(nf, "data_emissao", None)
        if not data_nf or not data_abertura:
            return None
        return max((data_abertura.date() - data_nf).days, 0)

    def importar_grupo(self, numero_origem, ano, grupo, validacao):
        primeira = grupo[0]
        row = primeira["row"]
        cols = primeira["cols"]
        numero_sac, sequencia = sac_numero_sistema(numero_origem, ano)
        cliente = validacao["cliente"]
        nf = validacao["nota"]
        concluido = validacao["concluido"]
        data_abertura = primeira["abertura"] or timezone.now()
        eventos = self._eventos_grupo(grupo)
        concluido_em = eventos.get("encerrado", {}).get("data") if concluido else None
        if concluido and not concluido_em:
            concluido_em = data_abertura

        titulo = f"SAC_{numero_sac} / {cliente.razao_social or cliente.nome} / NF {nf.numero_nf or nf.numero}"
        descricao = self.observacao_sac(row, cols)
        status_final = self.status_concluido if concluido else self.status_analise
        setor_final = None if concluido else validacao["setor"]
        acao_final = None if concluido else validacao["acao"]

        sac = SAC.objects.filter(empresa=self.empresa, numero=numero_sac).first()
        criado = sac is None

        if self.commit:
            defaults = {
                "ano": ano,
                "sequencia_ano": sequencia,
                "empresa": self.empresa,
                "cliente": cliente,
                "nota_fiscal": nf,
                "data_emissao_nf": nf.data_emissao,
                "numero_nf_revenda": limitar(valor_linha(row, cols["nf_cliente"]), 100) or None,
                "data_emissao_nf_revenda": data_valor(valor_linha(row, cols["dt_emissao_cliente"])),
                "tempo_uso_dias": self._tempo_uso_dias(nf, data_abertura),
                "status_inicial": "Aberto",
                "status_atual": status_final,
                "acao_em_espera": acao_final,
                "setor_responsavel": self.setores.get("sac"),
                "setor_atual": setor_final,
                "usuario_abertura": self.usuario,
                "contato_nome": limitar(valor_linha(row, cols["contato"]), 255),
                "telefone_1": limitar(valor_linha(row, cols["telefone"]), 30),
                "email_1": limitar(valor_linha(row, cols["email"]), 254),
                "titulo": limitar(titulo, 255),
                "descricao": descricao,
                "concluido_em": concluido_em,
                "email_automatico_habilitado": False,
            }
            if criado:
                sac = SAC.objects.create(numero=numero_sac, **defaults)
            else:
                for campo, valor in defaults.items():
                    setattr(sac, campo, valor)
                sac.save()
            SAC.objects.filter(pk=sac.pk).update(data_abertura=data_abertura)
            sac.data_abertura = data_abertura

        if criado:
            self.stats["sacs_criados"] += 1
        else:
            self.stats["sacs_atualizados"] += 1

        itens = self.importar_itens(sac, grupo, validacao) if self.commit else self.simular_itens(sac, grupo, validacao)
        if self.commit:
            self.registrar_historicos(sac, grupo, validacao, eventos, data_abertura, concluido_em)

        self.importados.append({"sac": numero_sac, "ano": ano, "linhas": [i["linha"] for i in grupo], "itens": itens})

    def simular_itens(self, sac, grupo, validacao):
        resultado = {"criados": 0, "atualizados": 0}
        sac_existente = sac or SAC.objects.filter(empresa=self.empresa, numero=validacao["numero_sac"]).first()
        for item in grupo:
            row = item["row"]
            cols = item["cols"]
            referencia = numero_texto(valor_linha(row, cols["referencia"]))
            equipamento = norm(valor_linha(row, cols["equipamento"]))
            item_nf = self._item_nf(validacao["nota"], referencia, equipamento)
            tipo_rastreio = "SERIAL" if norm_ascii(valor_linha(row, cols["categoria"])) == "equipamento" else "LOTE"
            numero_rastreio = numero_texto(valor_linha(row, cols["serie"]))
            existe = bool(
                sac_existente
                and item_nf
                and SACItem.objects.filter(
                    sac=sac_existente,
                    item_nota_fiscal=item_nf,
                    tipo_rastreio=tipo_rastreio,
                    numero_rastreio=numero_rastreio or "",
                ).exists()
            )
            if existe:
                resultado["atualizados"] += 1
                self.stats["itens_atualizados"] += 1
            else:
                resultado["criados"] += 1
                self.stats["itens_criados"] += 1
        return resultado

    def importar_itens(self, sac, grupo, validacao):
        resultado = {"criados": 0, "atualizados": 0}
        for item in grupo:
            row = item["row"]
            cols = item["cols"]
            referencia = numero_texto(valor_linha(row, cols["referencia"]))
            equipamento = norm(valor_linha(row, cols["equipamento"]))
            item_nf = self._item_nf(sac.nota_fiscal, referencia, equipamento)
            if not item_nf:
                continue
            tipo_rastreio = "SERIAL" if norm_ascii(valor_linha(row, cols["categoria"])) == "equipamento" else "LOTE"
            numero_rastreio = numero_texto(valor_linha(row, cols["serie"]))
            sac_item, criado = SACItem.objects.update_or_create(
                sac=sac,
                item_nota_fiscal=item_nf,
                tipo_rastreio=tipo_rastreio,
                numero_rastreio=numero_rastreio or "",
                defaults={
                    "tipo_ocorrencia": validacao["tipo"],
                    "quantidade_com_problema": inteiro(valor_linha(row, cols["quantidade"]), 1),
                    "grupo": limitar(valor_linha(row, cols["categoria"]), 255),
                    "item_quebrado_inservivel": None,
                    "item_pequeno_valor": None,
                    "observacao_item": self.observacao_item(row, cols),
                },
            )
            if criado:
                resultado["criados"] += 1
                self.stats["itens_criados"] += 1
            else:
                resultado["atualizados"] += 1
                self.stats["itens_atualizados"] += 1
        return resultado

    def registrar_historicos(self, sac, grupo, validacao, eventos, data_abertura, concluido_em):
        SacHistorico.objects.filter(
            sac=sac,
            acao_executada__in=[
                "Importacao historica EVENCOMERCIAL",
                "Historico legado EVENCOMERCIAL",
                "Conclusao historica EVENCOMERCIAL",
            ],
        ).delete()
        SacHistorico.objects.filter(sac=sac, acao_executada__startswith="Etapa historica EVENCOMERCIAL").delete()

        primeira = grupo[0]
        row = primeira["row"]
        cols = primeira["cols"]
        observacao_abertura = self.observacao_sac(row, cols) or "SAC importado da planilha historica Evencomercial."
        SacHistorico.objects.create(
            sac=sac,
            usuario=self.usuario,
            data_evento=data_abertura,
            status_anterior=None,
            status_novo=sac.status_atual,
            setor_origem=None,
            setor_destino=sac.setor_atual,
            acao_executada="Importacao historica EVENCOMERCIAL",
            observacao=observacao_abertura,
        )

        for evento in sorted(eventos.values(), key=lambda e: e["data"]):
            setor_destino = sac.setor_atual if evento["chave"] == "acao interna" and not validacao["concluido"] else None
            status_novo = self.status_concluido if evento["chave"] == "encerrado" else self.status_analise
            SacHistorico.objects.create(
                sac=sac,
                usuario=self.usuario,
                data_evento=evento["data"],
                status_anterior=None,
                status_novo=status_novo,
                setor_origem=None,
                setor_destino=setor_destino,
                acao_executada=f"Etapa historica EVENCOMERCIAL - {evento['label']}",
                observacao=evento["obs"],
            )

        historico = historico_formatado(valor_linha(row, cols["historico"]))
        if historico:
            SacHistorico.objects.create(
                sac=sac,
                usuario=self.usuario,
                data_evento=data_abertura,
                status_anterior=None,
                status_novo=sac.status_atual,
                setor_origem=None,
                setor_destino=sac.setor_atual,
                acao_executada="Historico legado EVENCOMERCIAL",
                observacao=historico,
            )

        if validacao["concluido"] and concluido_em:
            SacHistorico.objects.create(
                sac=sac,
                usuario=self.usuario,
                data_evento=concluido_em,
                status_anterior=self.status_aberto,
                status_novo=self.status_concluido,
                setor_origem=sac.setor_atual,
                setor_destino=sac.setor_atual,
                acao_executada="Conclusao historica EVENCOMERCIAL",
                observacao="SAC marcado como concluido conforme planilha historica Evencomercial.",
            )

    def _gerar_log_excel(self):
        if not self.log_output:
            return
        self.log_output.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Logs"
        headers = ["severidade", "sac", "linha", "campo", "valor", "motivo"]
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="0F4C75")
        white = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="B7D3E8")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = white
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for log in self.logs:
            ws.append([log.get(h, "") for h in headers])
        for row in ws.iter_rows(min_row=2):
            severidade = norm(row[0].value)
            fill = None
            if severidade == "ERRO":
                fill = PatternFill("solid", fgColor="F4CCCC")
            elif severidade == "AVISO":
                fill = PatternFill("solid", fgColor="FFF2CC")
            elif severidade == "IGNORADO":
                fill = PatternFill("solid", fgColor="EADCF8")
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
        for idx, header in enumerate(headers, 1):
            width = 18 if header not in {"valor", "motivo"} else 60
            ws.column_dimensions[get_column_letter(idx)].width = width

        resumo = wb.create_sheet("Resumo")
        resumo.append(["Indicador", "Valor"])
        for cell in resumo[1]:
            cell.fill = header_fill
            cell.font = white
            cell.border = border
        for key, value in [
            ("commit", self.commit),
            ("sacs_criados", self.stats["sacs_criados"]),
            ("sacs_atualizados", self.stats["sacs_atualizados"]),
            ("itens_criados", self.stats["itens_criados"]),
            ("itens_atualizados", self.stats["itens_atualizados"]),
            ("sacs_pulados", self.stats["sacs_pulados"]),
            ("logs_por_severidade", json.dumps(Counter(l["severidade"] for l in self.logs), ensure_ascii=False)),
        ]:
            resumo.append([key, value])
        resumo.column_dimensions["A"].width = 32
        resumo.column_dimensions["B"].width = 80
        wb.save(self.log_output)

    def executar(self):
        linhas = self.carregar()
        grupos = defaultdict(list)
        for item in linhas:
            grupos[item["chave"]].append(item)

        importacao = ImportacaoLegadoSAC.objects.create(
            nome=f"EVENCOMERCIAL SAC 2024/2026 - {self.arquivo.name}",
            arquivo_origem=str(self.arquivo),
            empresa=self.empresa,
            usuario=self.usuario,
            dry_run=not self.commit,
            status="COMMIT" if self.commit else "DRY_RUN",
            total_linhas=len(linhas),
            total_sacs_planilha=len(grupos),
        )

        try:
            with transaction.atomic():
                for (numero_origem, ano), grupo in sorted(grupos.items(), key=lambda x: (x[0][1], int(x[0][0]))):
                    validacao = self.validar_grupo(numero_origem, ano, grupo)
                    if validacao["avisos"] or validacao["pendencias"]:
                        self.stats["sacs_pulados"] += 1
                        if validacao["avisos"] and not validacao["pendencias"]:
                            self._log("AVISO", validacao["numero_sac"], grupo[0]["linha"], "SAC", "", "SAC com aviso ficou fora da importacao segura.")
                        continue
                    self.importar_grupo(numero_origem, ano, grupo, validacao)
                if not self.commit:
                    transaction.set_rollback(False)
        except Exception as exc:
            importacao.status = "ERRO"
            importacao.detalhes = {"erro": str(exc), "logs": self.logs, "importados": self.importados}
            importacao.save()
            raise

        importacao.sacs_criados = self.stats["sacs_criados"]
        importacao.sacs_atualizados = self.stats["sacs_atualizados"]
        importacao.itens_criados = self.stats["itens_criados"]
        importacao.itens_atualizados = self.stats["itens_atualizados"]
        importacao.sacs_pulados = self.stats["sacs_pulados"]
        importacao.detalhes = {
            "arquivo": str(self.arquivo),
            "commit": self.commit,
            "sacs_importados": self.importados,
            "logs": self.logs,
            "log_output": str(self.log_output or ""),
        }
        importacao.save()
        self._gerar_log_excel()
        return importacao


class Command(BaseCommand):
    help = "Importa SACs historicos da Evencomercial 2024/2026 com filtro seguro."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", help="Caminho da planilha EVEN_2024-2026.xlsx")
        parser.add_argument("--commit", action="store_true", help="Grava no banco. Sem esta flag, roda apenas simulacao.")
        parser.add_argument("--log-output", default="", help="Caminho para gerar log Excel.")
        parser.add_argument("--usuario", default="", help="Username responsavel pelo historico da importacao.")

    def handle(self, *args, **options):
        User = get_user_model()
        usuario = None
        username = (options.get("usuario") or "").strip()
        if username:
            usuario = User.objects.filter(username=username).first()
            if not usuario:
                raise CommandError(f"Usuario nao encontrado: {username}")
        if not usuario:
            usuario = User.objects.filter(is_superuser=True).first() or User.objects.first()
        if not usuario:
            raise CommandError("Nenhum usuario encontrado para registrar historico.")

        importador = ImportadorEvenComercialSAC(
            arquivo=options["arquivo"],
            commit=options["commit"],
            log_output=options.get("log_output") or "",
            usuario=usuario,
        )
        importacao = importador.executar()
        resumo = {
            "importacao_id": importacao.id,
            "status": importacao.status,
            "dry_run": importacao.dry_run,
            "total_linhas": importacao.total_linhas,
            "total_sacs_planilha": importacao.total_sacs_planilha,
            "sacs_criados": importacao.sacs_criados,
            "sacs_atualizados": importacao.sacs_atualizados,
            "itens_criados": importacao.itens_criados,
            "itens_atualizados": importacao.itens_atualizados,
            "sacs_pulados": importacao.sacs_pulados,
            "log_output": str(importador.log_output or ""),
        }
        self.stdout.write(json.dumps(resumo, ensure_ascii=False, indent=2))
