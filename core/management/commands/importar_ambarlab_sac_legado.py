from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import (
    AcaoEmEspera,
    Empresa,
    ImportacaoLegadoSAC,
    ItemNotaFiscal,
    NotaFiscal,
    SAC,
    SACItem,
    SacHistorico,
    Setor,
    StatusSAC,
    TipoOcorrencia,
)
from core.services.reincidencia_rastreio import sincronizar_historicos_reincidencia_sac
from core.services.texto_padrao import padronizar_caixa_texto


EMPRESA_CODIGO = "AMBARLAB"
MARCADOR_HISTORICO = "Importacao historica AMBARLAB"


CRONOLOGIA = [
    ("Aberto", "aberto", "SAC importado da planilha historica AMBARLAB."),
    (
        "Aberto - Aguardando Tratativa Tecnica Por Telefone",
        "aberto aguardando tratativa tecnica por telefone",
        "Etapa historica: aguardando tratativa tecnica por telefone.",
    ),
    ("Aberto - Aguardando Chegada", "aberto aguardando chegada", "Etapa historica: aguardando chegada."),
    ("Aberto - Aguardando Avaliação", "aberto aguardando avaliacao", "Etapa historica: aguardando avaliacao."),
    (
        "Aberto - Aguardando Aprov. Orçamento",
        "aberto aguardando aprov orcamento",
        "Etapa historica: aguardando aprovacao de orcamento.",
    ),
    ("Aberto - Aguardando Peça", "aberto aguardando peca", "Etapa historica: aguardando peca."),
    (
        "Aberto - Aguardando manutenção Interna",
        "aberto aguardando manutencao interna",
        "Etapa historica: aguardando manutencao interna.",
    ),
    (
        "Aberto - Aguardando manutenção Externa",
        "aberto aguardando manutencao externa",
        "Etapa historica: aguardando manutencao externa.",
    ),
    (
        "Aberto - Aguardando devolução ao Cliente",
        "aberto aguardando devolucao ao cliente",
        "Etapa historica: aguardando devolucao ao cliente.",
    ),
    (
        "Encerrado Parcialmente - Aguardando entrega ao cliente",
        "encerrado parcialmente aguardando entrega ao cliente",
        "Etapa historica: encerrado parcialmente, aguardando entrega ao cliente.",
    ),
    ("Ação interna", "acao interna", "Etapa historica: acao interna."),
    ("Encerrado", "encerrado", "SAC concluido conforme planilha historica AMBARLAB."),
]


COLUNAS = {
    "sac": "no sac",
    "status": "status",
    "aberto": "aberto",
    "encerrado": "encerrado",
    "acao_interna": "acao interna",
    "acao_padrao": "acao padrao",
    "agenda": "agenda previsao",
    "aguarda": "aguarda acao de",
    "qual_acao": "qual acao pendente",
    "nota": "nota fiscal venda",
    "cliente": "cliente",
    "contato": "contato",
    "telefone": "ddd telefone",
    "email": "e mail",
    "referencia": "referencia",
    "equipamento": "equipamento",
    "quantidade": "quantidade",
    "serie": "no serie",
    "relato": "o que o cliente alegou de problema",
    "parecer": "resumo parecer area tecnica interna terceira",
    "solucao": "solucao padrao",
    "tipo": "tipo de problema",
    "categoria": "categoria",
    "origem": "setor da origem do problema",
    "ato": "ato que gerou o problema",
    "garantia": "dentro do periodo de garantia",
    "mau_uso": "detectado mau uso",
    "conserto_garantia": "conserto em garantia",
}


def norm(valor):
    if valor in (None, ""):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return re.sub(r"\s+", " ", str(valor).strip().replace("\xa0", " "))


def norm_ascii(valor):
    texto = unicodedata.normalize("NFKD", norm(valor)).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[^a-zA-Z0-9]+", " ", texto.lower())
    return re.sub(r"\s+", " ", texto).strip()


def numero_texto(valor):
    texto = norm(valor)
    if not texto:
        return ""
    try:
        numero = float(texto.replace(",", "."))
        if numero.is_integer():
            return str(int(numero))
    except Exception:
        pass
    return texto


def inteiro(valor, padrao=1):
    texto = numero_texto(valor)
    if not texto:
        return padrao
    try:
        return max(1, int(Decimal(texto.replace(",", "."))))
    except (InvalidOperation, ValueError):
        return padrao


def limitar(valor, tamanho):
    texto = norm(valor)
    return texto[:tamanho] if texto else ""


def data_hora(valor, fim_do_dia=False):
    if not valor:
        return None
    if isinstance(valor, datetime):
        dt = valor
    else:
        return None
    if dt.time() == time.min and fim_do_dia:
        dt = datetime.combine(dt.date(), time.max.replace(microsecond=0))
    if timezone.is_naive(dt):
        return timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def sac_label(numero, ano):
    if str(numero).isdigit() and ano:
        return f"{int(numero):03d}/{ano}"
    return f"{numero}/{ano or '?'}"


def split_sac_label(label):
    numero, ano = label.split("/", 1)
    return int(numero), int(ano)


def historico_formatado(texto):
    bruto = norm(texto)
    if not bruto:
        return ""
    bruto = re.sub(r"\*{2,}", "\n", bruto)
    bruto = re.sub(r"(?<!^)(?<!\n)(?<!\d)(\d{1,2}/\d{1,2}/\d{2,4})\s*[-–:]", r"\n\1 -", bruto)
    linhas = []
    atual = ""
    for parte in re.split(r"\n+", bruto):
        parte = parte.strip(" -*")
        if not parte:
            continue
        parte = re.sub(r"^(\d{1,2}/\d{1,2}/\d{2,4})\s*[-–:]?\s*", r"\1 - ", parte)
        parte = re.sub(r"\s+-\s+", " - ", parte)
        if re.match(r"^\d{1,2}/\d{1,2}/\d{2,4}\s+-\s+", parte):
            if atual:
                linhas.append(atual.strip())
            atual = parte
        elif atual:
            atual += " " + parte
        else:
            linhas.append(parte)
    if atual:
        linhas.append(atual.strip())
    return "\n".join(linhas)


class ImportadorAmbarlabSAC:
    def __init__(self, arquivo, analise, commit, log_output, usuario):
        self.arquivo = Path(arquivo)
        self.analise_path = Path(analise)
        self.commit = commit
        self.log_output = Path(log_output) if log_output else None
        self.usuario = usuario
        self.logs = []
        self.resumo = []
        self.stats = Counter()
        self.sacs_importados = []

        self.empresa = Empresa.objects.get(codigo=EMPRESA_CODIGO)
        self.status_aberto = StatusSAC.objects.get(codigo="ABERTO")
        self.status_analise = StatusSAC.objects.get(codigo="EM_ANALISE")
        self.status_concluido = StatusSAC.objects.get(codigo="CONCLUIDO")
        self.setores = {norm_ascii(s.nome): s for s in Setor.objects.all()}
        self.acoes = {norm_ascii(a.nome): a for a in AcaoEmEspera.objects.all()}
        self.tipos = {norm_ascii(t.nome): t for t in TipoOcorrencia.objects.all()}

    def _log(self, tipo, sac, linha, campo, valor, mensagem):
        self.logs.append(
            {
                "tipo": tipo,
                "sac": sac,
                "linha": linha,
                "campo": campo,
                "valor": valor,
                "mensagem": mensagem,
            }
        )

    def _nota(self, numero):
        numero = numero_texto(numero)
        if not numero:
            return None
        return (
            NotaFiscal.objects.filter(empresa=self.empresa, numero_nf=numero).first()
            or NotaFiscal.objects.filter(empresa=self.empresa, numero=numero).first()
        )

    def _item_nf(self, nf, referencia, descricao):
        if not nf:
            return None
        qs = ItemNotaFiscal.objects.filter(nota_fiscal=nf)
        if referencia:
            item = qs.filter(codigo_produto__iexact=referencia).first()
            if item:
                return item
        if descricao:
            item = qs.filter(descricao_item__iexact=descricao).first()
            if item:
                return item
            desc_norm = norm_ascii(descricao)
            for candidato in qs:
                if norm_ascii(candidato.descricao_item) == desc_norm:
                    return candidato
        return None

    def _carregar_analise(self):
        dados = json.loads(self.analise_path.read_text(encoding="utf-8"))
        sacs = {row["SAC sistema"]: row for row in dados["sacs"] if row["Situacao"] != "BLOQUEADO"}
        itens = defaultdict(dict)
        for item in dados["itens"]:
            sac = item["SAC sistema"]
            if sac not in sacs:
                continue
            itens[sac][str(item.get("Linha") or "")] = item
        return dados, sacs, itens

    def _colunas(self, headers):
        idx = {norm_ascii(h): pos for pos, h in enumerate(headers) if norm(h)}

        def achar(chave):
            alvo = COLUNAS[chave]
            if alvo in idx:
                return idx[alvo]
            for nome, pos in idx.items():
                if alvo and alvo in nome:
                    return pos
            return None

        cols = {chave: achar(chave) for chave in COLUNAS}
        cols["cronologia"] = {chave: idx.get(chave) for _, chave, _ in CRONOLOGIA}
        return cols

    def _valor(self, row, col):
        if col is None or col >= len(row):
            return None
        return row[col]

    def _carregar_origem(self, sacs_liberados, itens_liberados):
        wb = openpyxl.load_workbook(self.arquivo, read_only=True, data_only=True)
        grupos = defaultdict(list)
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            headers = list(next(rows))
            cols = self._colunas(headers)
            for linha, row in enumerate(rows, 2):
                if not any(v not in (None, "") for v in row):
                    continue
                numero = numero_texto(self._valor(row, cols["sac"]))
                aberto = data_hora(self._valor(row, cols["aberto"]))
                if not numero and ws.title == "2025-2026" and linha == 5:
                    numero = "201"
                if not numero or not aberto:
                    continue
                label = sac_label(numero, aberto.year)
                if label not in sacs_liberados:
                    continue
                if str(linha) not in itens_liberados.get(label, {}) and any(
                    norm(self._valor(row, cols[campo])) for campo in ("referencia", "equipamento", "tipo")
                ):
                    continue
                grupos[label].append(
                    {
                        "aba": ws.title,
                        "linha": linha,
                        "row": row,
                        "cols": cols,
                        "item_analise": itens_liberados.get(label, {}).get(str(linha)),
                    }
                )
        return grupos

    def _eventos(self, grupo):
        eventos = {}
        for item in grupo:
            row = item["row"]
            cols = item["cols"]
            for label, chave, obs in CRONOLOGIA:
                data = data_hora(self._valor(row, cols["cronologia"].get(chave)), fim_do_dia=(chave == "encerrado"))
                if not data:
                    continue
                eventos[(data, label)] = {"data": data, "label": label, "obs": obs}
        return sorted(eventos.values(), key=lambda e: (e["data"], e["label"]))

    def _descricao_sac(self, grupo):
        blocos = []
        for item in grupo:
            row = item["row"]
            cols = item["cols"]
            partes = []
            relato = norm(self._valor(row, cols["relato"]))
            if relato:
                partes.append("Relato do cliente/usuário sobre o problema:\n" + padronizar_caixa_texto(relato))
            qual_acao = historico_formatado(self._valor(row, cols["qual_acao"]))
            if qual_acao:
                partes.append("Histórico recente:\n" + padronizar_caixa_texto(qual_acao))
            parecer = historico_formatado(self._valor(row, cols["parecer"]))
            if parecer:
                partes.append("Parecer técnico:\n" + padronizar_caixa_texto(parecer))
            solucao = norm(self._valor(row, cols["solucao"]))
            if solucao:
                partes.append("Solução padrão:\n" + padronizar_caixa_texto(solucao))
            if partes:
                blocos.append("\n\n".join(partes))
        return "\n\n---\n\n".join(dict.fromkeys(blocos)) or "SAC importado da planilha historica AMBARLAB."

    def _observacao_item(self, item):
        row = item["row"]
        cols = item["cols"]
        linhas = []
        for rotulo, campo in [
            ("Categoria", "categoria"),
            ("Setor da origem", "origem"),
            ("Ato que gerou o problema", "ato"),
            ("Garantia", "garantia"),
            ("Mau uso", "mau_uso"),
            ("Conserto em garantia", "conserto_garantia"),
        ]:
            valor = norm(self._valor(row, cols[campo]))
            if valor:
                linhas.append(f"{rotulo}: {valor}")
        return "\n".join(linhas)

    def _validar_sac(self, sac_row, grupo):
        bloqueios = []
        nf = self._nota(sac_row.get("NF origem"))
        if not nf:
            bloqueios.append(f"NF nao encontrada: {sac_row.get('NF origem')}")
        status_sugerido = norm_ascii(sac_row.get("Status sistema sugerido"))
        concluido = status_sugerido in {"concluido", "concluida"}
        setor = None
        acao = None
        if not concluido:
            setor = self.setores.get(norm_ascii(sac_row.get("Setor atual")))
            acao = self.acoes.get(norm_ascii(sac_row.get("Acao em espera")))
            if not setor and acao and acao.setor_destino:
                setor = acao.setor_destino
            if not setor:
                bloqueios.append("Setor atual nao encontrado")
            if not acao:
                bloqueios.append("Acao em espera nao encontrada")
        eventos = self._eventos(grupo)
        abertura = min([e["data"] for e in eventos if e["label"] == "Aberto"] or [None])
        if not abertura:
            bloqueios.append("Data de abertura nao encontrada")
        return {
            "bloqueios": bloqueios,
            "nota": nf,
            "concluido": concluido,
            "setor": setor,
            "acao": acao,
            "eventos": eventos,
            "abertura": abertura,
            "conclusao": max([e["data"] for e in eventos if e["label"] == "Encerrado"] or [None]) if concluido else None,
        }

    def _criar_ou_atualizar_sac(self, label, sac_row, grupo, validacao):
        sequencia, ano = split_sac_label(label)
        nf = validacao["nota"]
        cliente = nf.cliente
        primeiro = grupo[0]
        row = primeiro["row"]
        cols = primeiro["cols"]
        sac = SAC.objects.filter(empresa=self.empresa, numero=label).first()
        criado = sac is None
        titulo = f"SAC_{label} / {cliente.razao_social or cliente.nome} / NF {nf.numero_nf or nf.numero}"
        defaults = {
            "ano": ano,
            "sequencia_ano": sequencia,
            "empresa": self.empresa,
            "cliente": cliente,
            "nota_fiscal": nf,
            "data_emissao_nf": nf.data_emissao,
            "status_inicial": "Aberto",
            "status_atual": self.status_concluido if validacao["concluido"] else self.status_analise,
            "acao_em_espera": None if validacao["concluido"] else validacao["acao"],
            "setor_responsavel": self.setores.get("sac"),
            "setor_atual": None if validacao["concluido"] else validacao["setor"],
            "usuario_abertura": self.usuario,
            "contato_nome": limitar(self._valor(row, cols["contato"]), 255),
            "telefone_1": limitar(self._valor(row, cols["telefone"]), 30),
            "email_1": limitar(self._valor(row, cols["email"]), 254),
            "titulo": limitar(titulo, 255),
            "descricao": self._descricao_sac(grupo),
            "concluido_em": validacao["conclusao"],
            "email_automatico_habilitado": False,
        }
        if criado:
            sac = SAC.objects.create(numero=label, **defaults)
        else:
            for campo, valor in defaults.items():
                setattr(sac, campo, valor)
            sac.save()
        SAC.objects.filter(pk=sac.pk).update(data_abertura=validacao["abertura"])
        sac.data_abertura = validacao["abertura"]
        self.stats["sacs_criados" if criado else "sacs_atualizados"] += 1
        return sac, criado

    def _importar_itens(self, sac, grupo, validacao):
        resultado = Counter()
        for item in grupo:
            item_analise = item.get("item_analise")
            if not item_analise:
                continue
            referencia = numero_texto(item_analise.get("Referencia"))
            descricao = norm(item_analise.get("Descricao planilha"))
            if not referencia and not descricao:
                resultado["ignorados"] += 1
                self.stats["itens_ignorados"] += 1
                continue
            nf_item = self._nota(item_analise.get("NF usada")) or validacao["nota"]
            item_nf = self._item_nf(nf_item, referencia, descricao)
            if not item_nf:
                resultado["ignorados"] += 1
                self.stats["itens_ignorados"] += 1
                self._log("AVISO", sac.numero, item["linha"], "item", referencia or descricao, "Item da NF nao encontrado; item ignorado.")
                continue
            tipo = self.tipos.get(norm_ascii(item_analise.get("Tipo ocorrencia PARA")))
            categoria = norm(item_analise.get("Categoria"))
            serie = numero_texto(item_analise.get("Serie/Lote"))
            tipo_rastreio = "SERIAL" if norm_ascii(categoria) == "equipamento" or serie not in {"", "S/N", "SN"} else "LOTE"
            quantidade = inteiro(item_analise.get("Quantidade problema"), 1)
            observacao_item = self._observacao_item(item)
            numero_rastreio = limitar(serie, 150)
            if serie and len(serie) > 150:
                observacao_item = (observacao_item + "\n" if observacao_item else "") + f"Rastreio completo legado: {serie}"
                self._log(
                    "AVISO",
                    sac.numero,
                    item["linha"],
                    "numero_rastreio",
                    serie,
                    "Numero de rastreio maior que 150 caracteres; valor foi truncado no campo e preservado na observacao.",
                )
            _, criado = SACItem.objects.update_or_create(
                sac=sac,
                item_nota_fiscal=item_nf,
                tipo_rastreio=tipo_rastreio,
                numero_rastreio=numero_rastreio or "",
                defaults={
                    "tipo_ocorrencia": tipo,
                    "quantidade_com_problema": quantidade,
                    "grupo": limitar(categoria, 255),
                    "item_quebrado_inservivel": None,
                    "item_pequeno_valor": None,
                    "observacao_item": observacao_item,
                },
            )
            resultado["criados" if criado else "atualizados"] += 1
            self.stats["itens_criados" if criado else "itens_atualizados"] += 1
        return dict(resultado)

    def _registrar_historicos(self, sac, grupo, validacao):
        SacHistorico.objects.filter(sac=sac, acao_executada__startswith=MARCADOR_HISTORICO).delete()
        status_anterior = None
        setor_anterior = None
        for evento in validacao["eventos"]:
            status_novo = self.status_concluido if evento["label"] == "Encerrado" else self.status_analise
            setor_destino = None if evento["label"] == "Encerrado" else sac.setor_atual
            SacHistorico.objects.create(
                sac=sac,
                usuario=self.usuario,
                data_evento=evento["data"],
                status_anterior=status_anterior,
                status_novo=status_novo,
                setor_origem=setor_anterior,
                setor_destino=setor_destino,
                acao_executada=f"{MARCADOR_HISTORICO} - {evento['label']}",
                observacao=evento["obs"],
            )
            status_anterior = status_novo
            setor_anterior = setor_destino
        SacHistorico.objects.create(
            sac=sac,
            usuario=self.usuario,
            data_evento=validacao["abertura"],
            status_anterior=None,
            status_novo=sac.status_atual,
            setor_origem=None,
            setor_destino=sac.setor_atual,
            acao_executada=f"{MARCADOR_HISTORICO} - Histórico recente legado",
            observacao=sac.descricao,
        )

    def _salvar_log_excel(self):
        if not self.log_output:
            return
        self.log_output.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        headers = ["sac", "resultado", "detalhe", "itens"]
        ws.append(headers)
        for linha in self.resumo:
            ws.append([linha.get(h, "") for h in headers])
        log_ws = wb.create_sheet("Logs")
        log_headers = ["tipo", "sac", "linha", "campo", "valor", "mensagem"]
        log_ws.append(log_headers)
        for linha in self.logs:
            log_ws.append([linha.get(h, "") for h in log_headers])
        fill = PatternFill("solid", fgColor="0B4778")
        border = Border(left=Side(style="thin", color="B8D4E8"), right=Side(style="thin", color="B8D4E8"), top=Side(style="thin", color="B8D4E8"), bottom=Side(style="thin", color="B8D4E8"))
        for sheet in wb.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = fill
            for row in sheet.iter_rows():
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for col in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 28
        wb.save(self.log_output)

    def executar(self):
        _, sacs_liberados, itens_liberados = self._carregar_analise()
        grupos = self._carregar_origem(sacs_liberados, itens_liberados)
        importacao = ImportacaoLegadoSAC.objects.create(
            nome=f"AMBARLAB SAC legado - {'commit' if self.commit else 'simulacao'}",
            arquivo_origem=str(self.arquivo),
            empresa=self.empresa,
            usuario=self.usuario,
            dry_run=not self.commit,
            status="COMMIT" if self.commit else "DRY_RUN",
            total_linhas=sum(len(g) for g in grupos.values()),
            total_sacs_planilha=len(sacs_liberados),
        )
        try:
            with transaction.atomic():
                for label in sorted(sacs_liberados, key=lambda s: split_sac_label(s)):
                    grupo = grupos.get(label, [])
                    sac_row = sacs_liberados[label]
                    if not grupo:
                        self.stats["sacs_pulados"] += 1
                        self._log("BLOQUEIO", label, "", "origem", "", "SAC liberado no relatorio, mas nao encontrado na planilha origem.")
                        continue
                    validacao = self._validar_sac(sac_row, grupo)
                    if validacao["bloqueios"]:
                        self.stats["sacs_pulados"] += 1
                        detalhe = "; ".join(validacao["bloqueios"])
                        self._log("BLOQUEIO", label, "", "validacao", "", detalhe)
                        self.resumo.append({"sac": label, "resultado": "BLOQUEADO", "detalhe": detalhe, "itens": ""})
                        continue
                    if self.commit:
                        sac, criado = self._criar_ou_atualizar_sac(label, sac_row, grupo, validacao)
                        itens = self._importar_itens(sac, grupo, validacao)
                        self._registrar_historicos(sac, grupo, validacao)
                        self.sacs_importados.append(sac)
                    else:
                        criado = not SAC.objects.filter(empresa=self.empresa, numero=label).exists()
                        itens = {"simulados": len(itens_liberados.get(label, {}))}
                        self.stats["sacs_criados" if criado else "sacs_atualizados"] += 1
                        self.stats["itens_criados"] += itens["simulados"]
                    self.resumo.append({"sac": label, "resultado": "CRIAR" if criado else "ATUALIZAR", "detalhe": "", "itens": json.dumps(itens, ensure_ascii=False)})
                if not self.commit:
                    transaction.set_rollback(False)
        except Exception as exc:
            importacao.status = "ERRO"
            importacao.detalhes = {"erro": str(exc), "logs": self.logs[:500]}
            importacao.save()
            raise

        if self.commit and self.sacs_importados:
            reincidencia = sincronizar_historicos_reincidencia_sac(sacs=self.sacs_importados, usuario=self.usuario)
        else:
            reincidencia = {}
        importacao.sacs_criados = self.stats["sacs_criados"]
        importacao.sacs_atualizados = self.stats["sacs_atualizados"]
        importacao.itens_criados = self.stats["itens_criados"]
        importacao.itens_atualizados = self.stats["itens_atualizados"]
        importacao.itens_ignorados = self.stats["itens_ignorados"]
        importacao.sacs_pulados = self.stats["sacs_pulados"]
        importacao.detalhes = {
            "commit": self.commit,
            "logs": self.logs[:500],
            "log_output": str(self.log_output or ""),
            "reincidencia": reincidencia,
        }
        importacao.save()
        self._salvar_log_excel()
        return importacao


class Command(BaseCommand):
    help = "Importa SACs historicos da AMBARLAB usando relatorio pre-importacao aprovado."

    def add_arguments(self, parser):
        parser.add_argument("arquivo", help="Planilha AMBARLAB_2023-2026.xlsx")
        parser.add_argument("--analise", required=True, help="JSON final da analise pre-importacao")
        parser.add_argument("--commit", action="store_true", help="Grava no banco. Sem isso roda simulacao.")
        parser.add_argument("--log-output", default="", help="Caminho do Excel de log")

    def handle(self, *args, **options):
        usuario = get_user_model().objects.order_by("id").first()
        if not usuario:
            raise CommandError("Nenhum usuario encontrado para registrar historico.")
        importador = ImportadorAmbarlabSAC(
            arquivo=options["arquivo"],
            analise=options["analise"],
            commit=options["commit"],
            log_output=options.get("log_output") or "",
            usuario=usuario,
        )
        importacao = importador.executar()
        self.stdout.write(
            json.dumps(
                {
                    "importacao_id": importacao.id,
                    "status": importacao.status,
                    "dry_run": importacao.dry_run,
                    "total_linhas": importacao.total_linhas,
                    "total_sacs_planilha": importacao.total_sacs_planilha,
                    "sacs_criados": importacao.sacs_criados,
                    "sacs_atualizados": importacao.sacs_atualizados,
                    "itens_criados": importacao.itens_criados,
                    "itens_atualizados": importacao.itens_atualizados,
                    "itens_ignorados": importacao.itens_ignorados,
                    "sacs_pulados": importacao.sacs_pulados,
                    "log_output": str(importador.log_output or ""),
                },
                ensure_ascii=False,
                indent=2,
            )
        )
